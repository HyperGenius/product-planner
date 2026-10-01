# __tests__/unit/services/test_daily_report_parser.py
"""日報Excelパーサーの単体テスト (Issue #487)。

フィクスチャはすべてダミー値（実在の担当者名・顧客名等は使わない）。
"""

from datetime import date, datetime
from io import BytesIO
from typing import Any

import pytest
from app.services.daily_report_parser import (
    DailyReportParseError,
    parse_daily_report_workbook,
    parse_sheet_rows,
    parse_work_date,
    sheet_period,
)
from openpyxl import Workbook

HEADER = [
    "加工日",
    "担当者",
    "顧客先",
    "内職者名",
    "商品名",
    "工程名",
    "使用設備No",
    "加工数",
    "不適合合計数",
    "不良率",
    "修正内容",
    "修正数",
    "修正率",
    "不適合内容",
    "不適合数",
    "不適合内容",
    "不適合数",
    "段取り数",
    "ロットNo",
    "4M変更などを記載する備考欄",
]
COL = {
    name: i for i, name in enumerate(HEADER) if name not in ("不適合内容", "不適合数")
}
DEFECT_ITEM_COLS = [14, 16]


def _row(**values) -> list:
    """ヘッダ名をキーにした値から1行を作る。不適合数は defects=[n, m] で渡す。"""
    row: list = [None] * len(HEADER)
    defects = values.pop("defects", [])
    for name, value in values.items():
        row[COL[name]] = value
    for col, qty in zip(DEFECT_ITEM_COLS, defects, strict=False):
        row[col - 1] = "キズ"
        row[col] = qty
    return row


def _sheet_rows(*data_rows: list) -> list[list]:
    """実際の日報と同じく、上部に集計値・注意書きの行を置いた4行目にヘッダを置く。"""
    return [
        [100, 3, 0.03],
        ["各項目をクリックすると選択肢が表示されます。"],
        ["不適合合計数には検査式が入っています。"],
        HEADER,
        [],
        *data_rows,
    ]


def _parse(*data_rows: list, sheet_name: str = "2609製造"):
    return parse_sheet_rows(sheet_name, _sheet_rows(*data_rows))


def _codes(entry: dict, field: str) -> list[str]:
    return [i["code"] for i in entry["parse_issues"] if i["field"] == field]


class TestSheetPeriod:
    def test_target_sheet(self):
        assert sheet_period("2609製造") == (2026, 9)

    @pytest.mark.parametrize(
        "name",
        ["原紙", "リスト", "設備台帳目次", "2613製造", "26095製造", "2609製造 (2)"],
    )
    def test_non_target_sheets(self, name):
        assert sheet_period(name) is None


class TestParseWorkDate:
    def test_normal_date(self):
        assert parse_work_date("2026.09.01", (2026, 9)) == (date(2026, 9, 1), [])

    def test_year_typo_is_corrected_to_sheet_year(self):
        assert parse_work_date("2029.09.02", (2026, 9)) == (
            date(2026, 9, 2),
            ["year_corrected"],
        )

    def test_concatenated_dates_use_first(self):
        assert parse_work_date("2026.09.232026.09.18", (2026, 9)) == (
            date(2026, 9, 23),
            ["concatenated"],
        )

    def test_previous_month_is_kept(self):
        assert parse_work_date("2026.08.31", (2026, 9)) == (
            date(2026, 8, 31),
            ["outside_sheet_month"],
        )

    def test_previous_month_across_year_is_not_year_corrected(self):
        # 2026年1月のシートの前年12月の行は、年の誤記ではなく前月の行
        assert parse_work_date("2025.12.31", (2026, 1)) == (
            date(2025, 12, 31),
            ["outside_sheet_month"],
        )

    def test_year_typo_in_january_sheet_for_december_row(self):
        assert parse_work_date("2052.12.28", (2026, 1)) == (
            date(2025, 12, 28),
            ["year_corrected", "outside_sheet_month"],
        )

    def test_far_date_is_kept_with_issue(self):
        assert parse_work_date("2026.03.05", (2026, 9)) == (
            date(2026, 3, 5),
            ["outside_sheet_month"],
        )

    def test_full_width_digits(self):
        assert parse_work_date("２０２６．０９．０３", (2026, 9)) == (
            date(2026, 9, 3),
            [],
        )

    def test_datetime_cell(self):
        assert parse_work_date(datetime(2026, 9, 4), (2026, 9)) == (
            date(2026, 9, 4),
            [],
        )

    @pytest.mark.parametrize(
        "raw", ["9/1", "2026.09.01 午前", "2026.07.06 メモ", "2026.13.01"]
    )
    def test_invalid(self, raw):
        work_date, issues = parse_work_date(raw, (2026, 9))
        assert work_date is None
        assert issues[-1] == "invalid"

    @pytest.mark.parametrize("raw", [None, "", "  "])
    def test_missing(self, raw):
        assert parse_work_date(raw, (2026, 9)) == (None, ["missing"])


class TestParseSheetRows:
    def test_basic_row(self):
        sheet = _parse(
            _row(
                加工日="2026.09.01",
                担当者=" 作業者A ",
                顧客先="顧客A社",
                内職者名="内職者X",
                商品名="製品A-10φｘ8L",
                工程名="カシメ加工",
                使用設備No="プレス15t 3号機",
                加工数=1000,
                不適合合計数=3,
                段取り数=4,
                ロットNo="L-01",
                **{"4M変更などを記載する備考欄": "金型交換"},
                defects=[1, 2],
            )
        )
        assert sheet.header_found
        assert sheet.entries == [
            {
                "row_no": 6,
                "work_date": "2026-09-01",
                "work_date_raw": "2026.09.01",
                "worker_raw": "作業者A",
                "customer_raw": "顧客A社",
                "homeworker_raw": "内職者X",
                "product_raw": "製品A-10φｘ8L",
                "process_raw": "カシメ加工",
                "equipment_raw": "プレス15t 3号機",
                "processed_qty": 1000,
                "defect_qty": 3,
                "good_qty": 997,
                "setup_qty": 4,
                "lot_no": "L-01",
                "note": "金型交換",
                "parse_issues": [],
            }
        ]

    def test_header_row_is_located_by_cell_not_fixed_row(self):
        rows = (
            [["注意書き"]] * 7
            + [HEADER]
            + [_row(加工日="2026.09.01", 商品名="P", 加工数=5)]
        )
        sheet = parse_sheet_rows("2609製造", rows)
        assert [e["row_no"] for e in sheet.entries] == [9]

    def test_columns_are_resolved_by_header_name(self):
        # 列の並びが変わってもヘッダ名で引く
        rows: list[list[Any]] = [
            ["商品名", "加工数", "加工日", "担当者"],
            ["P-1", 10, "2026.09.01", "作業者A"],
        ]
        (entry,) = parse_sheet_rows("2609製造", rows).entries
        assert entry["product_raw"] == "P-1"
        assert entry["processed_qty"] == 10
        assert entry["work_date"] == "2026-09-01"
        assert entry["worker_raw"] == "作業者A"
        assert entry["defect_qty"] == 0
        assert entry["good_qty"] == 10

    def test_no_header(self):
        sheet = parse_sheet_rows("2609製造", [["メモ"], ["2026.09.01", "P", 1]])
        assert sheet.header_found is False
        assert sheet.entries == []

    def test_rejects_non_target_sheet(self):
        with pytest.raises(ValueError):
            parse_sheet_rows("原紙", [HEADER])

    def test_defect_total_falls_back_to_item_sum(self):
        (entry,) = _parse(
            _row(加工日="2026.09.01", 商品名="P", 加工数=100, defects=[2, 5])
        ).entries
        assert entry["defect_qty"] == 7
        assert entry["good_qty"] == 93

    def test_good_qty_is_clamped_to_zero(self):
        (entry,) = _parse(
            _row(加工日="2026.09.01", 商品名="P", 加工数=3, 不適合合計数=5)
        ).entries
        assert entry["good_qty"] == 0
        assert _codes(entry, "good_qty") == ["negative_clamped"]

    def test_missing_processed_qty_keeps_row_with_null_good_qty(self):
        (entry,) = _parse(_row(加工日="2026.09.01", 商品名="P", 不適合合計数=0)).entries
        assert entry["processed_qty"] is None
        assert entry["good_qty"] is None
        assert _codes(entry, "processed_qty") == ["missing"]

    def test_invalid_processed_qty(self):
        (entry,) = _parse(_row(加工日="2026.09.01", 商品名="P", 加工数="約100")).entries
        assert entry["processed_qty"] is None
        assert entry["good_qty"] is None
        assert _codes(entry, "processed_qty") == ["invalid"]

    def test_quantity_strings_are_normalized(self):
        (entry,) = _parse(
            _row(加工日="2026.09.01", 商品名="P", 加工数="１,２００", 不適合合計数=1.0)
        ).entries
        assert entry["processed_qty"] == 1200
        assert entry["defect_qty"] == 1
        assert entry["good_qty"] == 1199

    def test_numeric_product_name_is_stringified(self):
        entries = _parse(
            _row(加工日="2026.09.01", 商品名=1234, 加工数=1),
            _row(加工日="2026.09.01", 商品名=5678.0, 加工数=1),
        ).entries
        assert [e["product_raw"] for e in entries] == ["1234", "5678"]

    def test_multiple_workers_are_kept_raw(self):
        entries = _parse(
            _row(加工日="2026.09.01", 担当者="作業者A・作業者B", 商品名="P", 加工数=1),
            _row(加工日="2026.09.01", 担当者="作業者A/作業者B", 商品名="P", 加工数=1),
        ).entries
        assert [e["worker_raw"] for e in entries] == [
            "作業者A・作業者B",
            "作業者A/作業者B",
        ]

    def test_date_anomalies_are_recorded_without_dropping_rows(self):
        entries = _parse(
            _row(加工日="2026.08.31", 商品名="P", 加工数=1),
            _row(加工日="2029.09.02", 商品名="P", 加工数=1),
            _row(加工日="2026.09.232026.09.18", 商品名="P", 加工数=1),
            _row(加工日="不明", 商品名="P", 加工数=1),
        ).entries
        assert [(e["work_date"], _codes(e, "work_date")) for e in entries] == [
            ("2026-08-31", ["outside_sheet_month"]),
            ("2026-09-02", ["year_corrected"]),
            ("2026-09-23", ["concatenated"]),
            (None, ["invalid"]),
        ]
        assert entries[1]["work_date_raw"] == "2029.09.02"

    def test_memo_and_blank_rows_are_skipped(self):
        entries = _parse(
            _row(加工日="2026.09.01", 商品名="P", 加工数=1),
            [],
            [None] * len(HEADER),
            _row(加工日="3.24 製品B 100個 備考欄に記載無し"),
            _row(加工日="2026.07.06 シートの書式を変更した。"),
            # 不適合合計数の数式のキャッシュ値（0）だけが残っている行
            _row(不適合合計数=0),
            # 加工日だけ先に入れた行
            _row(加工日="2026.09.30"),
        ).entries
        assert [e["row_no"] for e in entries] == [6]

    def test_row_with_unparsable_date_but_product_is_kept(self):
        (entry,) = _parse(_row(加工日="9/1", 商品名="P")).entries
        assert entry["work_date"] is None
        assert entry["work_date_raw"] == "9/1"


def _workbook_bytes(build) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    build(wb)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestParseDailyReportWorkbook:
    def test_reads_only_target_sheets(self):
        def build(wb):
            for title in ["設備台帳目次", "原紙", "2608製造", "2609製造", "リスト"]:
                ws = wb.create_sheet(title)
                for row in _sheet_rows(_row(加工日="2026.09.01", 商品名="P", 加工数=1)):
                    ws.append(row)

        sheets = parse_daily_report_workbook(_workbook_bytes(build))
        assert [s.sheet_name for s in sheets] == ["2608製造", "2609製造"]
        assert all(len(s.entries) == 1 for s in sheets)

    def test_formula_without_cached_value_uses_item_sum(self):
        # openpyxl で保存した数式セルはキャッシュ値を持たない（data_only で None になる）
        def build(wb):
            ws = wb.create_sheet("2609製造")
            for row in _sheet_rows(
                _row(加工日="2026.09.01", 商品名="P", 加工数=50, defects=[1, 4])
            ):
                ws.append(row)
            ws.cell(row=6, column=COL["不適合合計数"] + 1, value="=O6+Q6")

        (sheet,) = parse_daily_report_workbook(_workbook_bytes(build))
        (entry,) = sheet.entries
        assert entry["defect_qty"] == 5
        assert entry["good_qty"] == 45

    @pytest.mark.parametrize(
        "content", [b"", b"not an excel file", b"PK\x03\x04broken"]
    )
    def test_unreadable_content(self, content):
        with pytest.raises(DailyReportParseError):
            parse_daily_report_workbook(content)
