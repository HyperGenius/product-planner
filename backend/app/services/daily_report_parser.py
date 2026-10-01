# backend/app/services/daily_report_parser.py
"""日報Excelのパーサー (Issue #487, 親Issue #485)。

Storage・DB から切り離した純粋関数。ブックのバイト列から `YYMM製造` シートの各行を
`daily_report_entries` に保存する形の dict にする。名寄せ（マスタとの照合）は行わず、
日報の値をそのまま（前後の空白除去・数値セルの文字列化のみ）持つ。

日報の書式（パイロット顧客の2026年9月分の調査結果）:
* シート: 月ごとの `YYMM製造` ＋ テンプレート `原紙` ＋ 入力規則の `リスト` 等
* シート上部に集計値・注意書きの行があり、その下に `加工日` から始まるヘッダ行がある
* 不適合は「不適合内容／不適合数」の組が複数列あり、`不適合合計数` はその合計の数式セル
* 加工日は `YYYY.MM.DD` の文字列。年の誤記・日付の連結・前月の行・シート下部のメモ行が混在する
"""

import calendar
import re
import unicodedata
import warnings
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

# 対象シート名。`原紙`（テンプレート）・`リスト`・`設備台帳目次` 等は読まない。
_TARGET_SHEET_PATTERN = re.compile(r"^(\d{2})(\d{2})製造$")
# ヘッダ行（`加工日` のセル）を探す範囲。上部の注意書きの行数は今後ずれうるので余裕を持たせる。
_HEADER_SEARCH_ROWS = 30
# `2026.09.01`。連結（`2026.09.232026.09.18`）はこの繰り返しとして扱う。
_DATE_TOKEN = r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})"
_DATE_PATTERN = re.compile(rf"^(?:{_DATE_TOKEN})+$")
_FIRST_DATE_PATTERN = re.compile(_DATE_TOKEN)
_INTEGER_PATTERN = re.compile(r"^-?\d+(?:\.0+)?$")

# 明細の列 → 日報のヘッダ名（NFKC・空白除去後に完全一致）
_HEADER_NAMES: dict[str, str] = {
    "work_date": "加工日",
    "worker": "担当者",
    "customer": "顧客先",
    "homeworker": "内職者名",
    "product": "商品名",
    "process": "工程名",
    "equipment": "使用設備No",
    "processed_qty": "加工数",
    "defect_total": "不適合合計数",
    "setup_qty": "段取り数",
    "lot_no": "ロットNo",
}
# 備考欄のヘッダは説明文（`4M変更+3H金型に…備考欄`）なので部分一致で引く
_NOTE_HEADER_KEYWORD = "備考"
# 個別の不適合数（複数列）。不適合合計数のキャッシュ値が無いときの代替に使う
_DEFECT_ITEM_HEADER = "不適合数"


class DailyReportParseError(Exception):
    """ブックとして読み込めない（壊れたファイル・Excel 以外のファイル）。"""


@dataclass
class ParsedSheet:
    sheet_name: str
    entries: list[dict[str, Any]] = field(default_factory=list)
    # ヘッダ行（`加工日`）が見つからなかった。この場合は明細を置き換えない
    header_found: bool = True


@dataclass(frozen=True)
class _SheetPeriod:
    year: int
    month: int

    @property
    def first_day(self) -> date:
        return date(self.year, self.month, 1)

    @property
    def last_day(self) -> date:
        return date(
            self.year, self.month, calendar.monthrange(self.year, self.month)[1]
        )

    def window(self) -> tuple[date, date]:
        """前月初〜翌月末。前月・翌月の行はシートの年月の誤記ではなく、そのまま採用する範囲。"""
        prev_year, prev_month = (
            (self.year - 1, 12) if self.month == 1 else (self.year, self.month - 1)
        )
        next_year, next_month = (
            (self.year + 1, 1) if self.month == 12 else (self.year, self.month + 1)
        )
        return (
            date(prev_year, prev_month, 1),
            date(next_year, next_month, calendar.monthrange(next_year, next_month)[1]),
        )


def sheet_period(sheet_name: str) -> tuple[int, int] | None:
    """`2609製造` → (2026, 9)。対象シートでなければ None。"""
    m = _TARGET_SHEET_PATTERN.fullmatch(sheet_name)
    if not m:
        return None
    month = int(m.group(2))
    if not 1 <= month <= 12:
        return None
    return 2000 + int(m.group(1)), month


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value)))


def _to_text(value: Any) -> str | None:
    """セルの値を raw の文字列にする。数値セル（商品名の `1234` 等）は文字列化する。"""
    if value is None:
        return None
    if isinstance(value, bool):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    elif isinstance(value, datetime):
        text = (
            value.date().isoformat()
            if value.time() == datetime.min.time()
            else value.isoformat()
        )
    elif isinstance(value, date):
        text = value.isoformat()
    else:
        text = str(value)
    text = text.strip()
    return text or None


def _to_int(value: Any) -> tuple[int | None, bool]:
    """数量セルを整数にする。(値, 解釈できたか)。空欄は (None, True)。"""
    if value is None:
        return None, True
    if isinstance(value, bool):
        return None, False
    if isinstance(value, int):
        return value, True
    if isinstance(value, float):
        return (int(value), True) if value.is_integer() else (None, False)
    text = unicodedata.normalize("NFKC", str(value)).strip().replace(",", "")
    if not text:
        return None, True
    if _INTEGER_PATTERN.fullmatch(text):
        return int(float(text)), True
    return None, False


def _safe_date(year: int, month: int, day: int) -> date | None:
    try:
        return date(year, month, day)
    except ValueError:
        return None


def parse_work_date(
    value: Any, period: tuple[int, int]
) -> tuple[date | None, list[str]]:
    """加工日を解釈する。(日付, issue コードの一覧)。解釈できなければ日付は None。

    * 連結（`2026.09.232026.09.18`）→ 先頭の日付を採用（`concatenated`）
    * シートの年月の前月〜翌月 → そのまま採用。シートの月以外なら `outside_sheet_month`
    * それ以外で、年をシートの年（年をまたぐ前月・翌月なら前後の年）に置き換えると前月〜翌月に
      収まる → 補正（`year_corrected`。例: 2026年9月のシートの `2029.09.02` → 2026-09-02）
    * 補正もできない日付は、そのまま採用して `outside_sheet_month`
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, ["missing"]
    parts, issues = _date_parts(value)
    if parts is None:
        return None, issues
    work_date, period_issues = _fit_to_sheet_period(parts, _SheetPeriod(*period))
    return work_date, issues + period_issues


def _date_parts(value: Any) -> tuple[tuple[int, int, int] | None, list[str]]:
    """加工日のセルを (年, 月, 日) にする。日付として読めなければ None と `invalid`。"""
    if isinstance(value, date):  # datetime も含む（日付セルとして入力された場合）
        return (value.year, value.month, value.day), []
    text = re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value)))
    if not _DATE_PATTERN.fullmatch(text):
        return None, ["invalid"]
    first = _FIRST_DATE_PATTERN.match(text)
    assert first is not None
    issues = ["concatenated"] if first.end() != len(text) else []
    return (int(first.group(1)), int(first.group(2)), int(first.group(3))), issues


def _fit_to_sheet_period(
    parts: tuple[int, int, int], sheet: _SheetPeriod
) -> tuple[date | None, list[str]]:
    """シートの年月（前月〜翌月）に照らして日付を確定する。年の誤記はここで補正する。"""
    window_start, window_end = sheet.window()

    def in_sheet_month(d: date) -> list[str]:
        return [] if sheet.first_day <= d <= sheet.last_day else ["outside_sheet_month"]

    year, month, day = parts
    parsed = _safe_date(year, month, day)
    if parsed is not None and window_start <= parsed <= window_end:
        return parsed, in_sheet_month(parsed)

    for candidate_year in (sheet.year, sheet.year - 1, sheet.year + 1):
        candidate = _safe_date(candidate_year, month, day)
        if (
            candidate_year != year
            and candidate is not None
            and window_start <= candidate <= window_end
        ):
            return candidate, ["year_corrected", *in_sheet_month(candidate)]

    if parsed is None:
        return None, ["invalid"]
    return parsed, ["outside_sheet_month"]


def _find_header(
    rows: Sequence[Sequence[Any]],
) -> tuple[int, dict[str, int], list[int]] | None:
    """(ヘッダ行の index, 列名→列 index, 個別の不適合数の列 index 一覧)。"""
    for row_index, row in enumerate(rows[:_HEADER_SEARCH_ROWS]):
        headers = [_normalize_header(v) for v in row]
        if _HEADER_NAMES["work_date"] not in headers:
            continue
        columns: dict[str, int] = {}
        for key, name in _HEADER_NAMES.items():
            if name in headers:
                columns[key] = headers.index(name)
        for col, header in enumerate(headers):
            if _NOTE_HEADER_KEYWORD in header:
                columns["note"] = col
                break
        defect_items = [
            col for col, header in enumerate(headers) if header == _DEFECT_ITEM_HEADER
        ]
        return row_index, columns, defect_items
    return None


def _cell(row: Sequence[Any], columns: dict[str, int], key: str) -> Any:
    col = columns.get(key)
    if col is None or col >= len(row):
        return None
    value = row[col]
    if isinstance(value, str) and not value.strip():
        return None
    return value


def _defect_qty(
    row: Sequence[Any], columns: dict[str, int], defect_items: list[int]
) -> tuple[int | None, bool]:
    """不適合合計数。数式セルのキャッシュ値が無ければ、個別の不適合数の合計で代替する。"""
    total, ok = _to_int(_cell(row, columns, "defect_total"))
    if total is not None:
        return total, True
    item_sum = 0
    for col in defect_items:
        qty, item_ok = _to_int(row[col] if col < len(row) else None)
        ok = ok and item_ok
        if qty is not None:
            item_sum += qty
    return item_sum, ok


def _parse_row(
    row: Sequence[Any],
    row_no: int,
    columns: dict[str, int],
    defect_items: list[int],
    period: tuple[int, int],
) -> dict[str, Any] | None:
    raw_date = _cell(row, columns, "work_date")
    text_fields = {
        "worker_raw": _to_text(_cell(row, columns, "worker")),
        "customer_raw": _to_text(_cell(row, columns, "customer")),
        "homeworker_raw": _to_text(_cell(row, columns, "homeworker")),
        "product_raw": _to_text(_cell(row, columns, "product")),
        "process_raw": _to_text(_cell(row, columns, "process")),
        "equipment_raw": _to_text(_cell(row, columns, "equipment")),
    }
    raw_processed = _cell(row, columns, "processed_qty")

    work_date, date_issues = parse_work_date(raw_date, period)
    has_product_or_qty = (
        text_fields["product_raw"] is not None or raw_processed is not None
    )
    # 日付として解釈できず商品名・加工数も空の行（シート下部のメモ行・空行）は明細にしない。
    # 加工日だけが入った行（入力途中・テンプレートの日付の先入れ）も同様。
    if not has_product_or_qty and (
        work_date is None or all(v is None for v in text_fields.values())
    ):
        return None

    issues = [{"field": "work_date", "code": code} for code in date_issues]

    processed_qty, processed_ok = _to_int(raw_processed)
    if not processed_ok:
        issues.append({"field": "processed_qty", "code": "invalid"})
    elif processed_qty is None:
        issues.append({"field": "processed_qty", "code": "missing"})

    defect_qty, defect_ok = _defect_qty(row, columns, defect_items)
    if not defect_ok:
        issues.append({"field": "defect_qty", "code": "invalid"})

    setup_qty, setup_ok = _to_int(_cell(row, columns, "setup_qty"))
    if not setup_ok:
        issues.append({"field": "setup_qty", "code": "invalid"})

    good_qty: int | None = None
    if processed_qty is not None:
        good_qty = processed_qty - (defect_qty or 0)
        if good_qty < 0:
            issues.append({"field": "good_qty", "code": "negative_clamped"})
            good_qty = 0

    return {
        "row_no": row_no,
        "work_date": work_date.isoformat() if work_date else None,
        "work_date_raw": _to_text(raw_date),
        **text_fields,
        "processed_qty": processed_qty,
        "defect_qty": defect_qty,
        "good_qty": good_qty,
        "setup_qty": setup_qty,
        "lot_no": _to_text(_cell(row, columns, "lot_no")),
        "note": _to_text(_cell(row, columns, "note")),
        "parse_issues": issues,
    }


def parse_sheet_rows(sheet_name: str, rows: Iterable[Sequence[Any]]) -> ParsedSheet:
    """1シート分のセルの値（`iter_rows(values_only=True)` の結果）を明細にする。

    `sheet_name` は `YYMM製造` であること（呼び出し側で `sheet_period()` を確認する）。
    `row_no` は Excel の行番号（1始まり）。
    """
    period = sheet_period(sheet_name)
    if period is None:
        raise ValueError(f"not a daily report sheet: {sheet_name}")

    all_rows = [tuple(r) for r in rows]
    header = _find_header(all_rows)
    if header is None:
        return ParsedSheet(sheet_name=sheet_name, header_found=False)
    header_index, columns, defect_items = header

    parsed = ParsedSheet(sheet_name=sheet_name)
    for row_index in range(header_index + 1, len(all_rows)):
        entry = _parse_row(
            all_rows[row_index], row_index + 1, columns, defect_items, period
        )
        if entry is not None:
            parsed.entries.append(entry)
    return parsed


def parse_daily_report_workbook(content: bytes) -> list[ParsedSheet]:
    """ブック中の `YYMM製造` シートをすべてパースする。

    数式セル（不適合合計数）はキャッシュ値を読む（`data_only=True`）。
    読み込めないファイルは `DailyReportParseError`。
    """
    try:
        with warnings.catch_warnings():
            # 入力規則の拡張（Data Validation extension）等、値の読み取りに関係ない警告を抑制する
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise DailyReportParseError("unable to open workbook") from exc

    try:
        sheets: list[ParsedSheet] = []
        for worksheet in workbook.worksheets:
            if sheet_period(worksheet.title) is None:
                continue
            # read_only モードはファイルに記録された使用範囲（dimensions）だけを読むため、
            # 作成元によって範囲が実データより狭いことがある。範囲を捨てて全行を読む。
            worksheet.reset_dimensions()
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                rows = list(worksheet.iter_rows(values_only=True))
            sheets.append(parse_sheet_rows(worksheet.title, rows))
        return sheets
    finally:
        workbook.close()
